import openpyxl

wb = openpyxl.load_workbook(r'c:\laragon\www\FindServices\keamanan_sistem_informasi\1. template_baru_Praktikum_STRIDE-DREAD.xlsx', data_only=False)

print("=== RINGKASAN FORMULAS ===")
ws_r = wb['Ringkasan']
for r in range(4, 25):
    b = ws_r[f'B{r}'].value
    c = ws_r[f'C{r}'].value
    print(f"Row {r:2d}: {b} -> {c}")

print("\n=== STRIDE_DREAD_Template HEADERS ===")
ws_t = wb['STRIDE_DREAD_Template']
for c in range(1, 23):
    col_let = openpyxl.utils.get_column_letter(c)
    val = ws_t[f'{col_let}4'].value
    val3 = ws_t[f'{col_let}3'].value
    print(f"{col_let}: row3={val3} | row4={val}")

print("\n=== STRIDE_DREAD_Template FORMULAS (Row 5) ===")
for c in range(1, 23):
    col_let = openpyxl.utils.get_column_letter(c)
    val = ws_t[f'{col_let}5'].value
    print(f"{col_let}5 = {val}")
