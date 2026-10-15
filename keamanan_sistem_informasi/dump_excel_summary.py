import openpyxl

wb = openpyxl.load_workbook(r'c:\laragon\www\FindServices\keamanan_sistem_informasi\1. template_baru_Praktikum_STRIDE-DREAD.xlsx', data_only=True)

print("=== SHEETS ===")
print(wb.sheetnames)

print("\n=== PROFIL APLIKASI (Section A) ===")
ws_p = wb['Profil_Aplikasi']
for r in range(5, 10):
    label = ws_p[f'A{r}'].value
    val = ws_p[f'B{r}'].value
    print(f"{label}: {val}")

print("\n=== KOMPONEN SISTEM (Section B) ===")
for r in range(15, 25):
    kid = ws_p[f'B{r}'].value
    knama = ws_p[f'C{r}'].value
    ktipe = ws_p[f'D{r}'].value
    kdesc = ws_p[f'E{r}'].value
    if kid:
        print(f"[{kid}] {knama} ({ktipe}): {kdesc}")

print("\n=== ALIRAN DATA (Section C) ===")
for r in range(30, 40):
    fid = ws_p[f'B{r}'].value
    fsrc = ws_p[f'C{r}'].value
    fdst = ws_p[f'D{r}'].value
    fdata = ws_p[f'E{r}'].value
    if fid:
        print(f"[{fid}] {fsrc} -> {fdst}: {fdata}")

print("\n=== TRUST BOUNDARIES (Section D) ===")
for r in range(45, 50):
    tbid = ws_p[f'B{r}'].value
    tbdesc = ws_p[f'C{r}'].value
    if tbid:
        print(f"[{tbid}]: {tbdesc}")

print("\n=== STRIDE & DREAD THREATS ===")
ws_t = wb['STRIDE_DREAD_Template']
for r in range(5, 45):
    tid = ws_t[f'B{r}'].value
    if not tid:
        continue
    desc = ws_t[f'C{r}'].value
    s_s = ws_t[f'D{r}'].value
    s_t = ws_t[f'E{r}'].value
    s_r = ws_t[f'F{r}'].value
    s_i = ws_t[f'G{r}'].value
    s_d = ws_t[f'H{r}'].value
    s_e = ws_t[f'I{r}'].value
    stride_cats = []
    if s_s: stride_cats.append('S')
    if s_t: stride_cats.append('T')
    if s_r: stride_cats.append('R')
    if s_i: stride_cats.append('I')
    if s_d: stride_cats.append('D')
    if s_e: stride_cats.append('E')
    
    d_d = ws_t[f'K{r}'].value
    d_r = ws_t[f'L{r}'].value
    d_e = ws_t[f'M{r}'].value
    d_a = ws_t[f'N{r}'].value
    d_x = ws_t[f'O{r}'].value
    total = ws_t[f'P{r}'].value
    lvl = ws_t[f'Q{r}'].value
    alasan = ws_t[f'S{r}'].value
    bukti = ws_t[f'T{r}'].value
    mitigasi = ws_t[f'U{r}'].value
    print(f"\n--- {tid} | {lvl} (Total: {total}) | STRIDE: {','.join(stride_cats)} ---")
    print(f"Deskripsi: {desc}")
    print(f"DREAD: D={d_d}, R={d_r}, E={d_e}, A={d_a}, D={d_x}")
    print(f"Alasan: {alasan}")
    print(f"Bukti: {bukti}")
    print(f"Mitigasi: {mitigasi}")

print("\n=== RINGKASAN SHEET ===")
ws_r = wb['Ringkasan']
for r in range(1, 35):
    vals = [str(ws_r.cell(r, c).value or '') for c in range(1, 8)]
    if any(v.strip() for v in vals):
        print(f"Row {r:2d}: " + " | ".join(vals))
